# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx", "openpyxl", "pypdf"]
# ///
"""Gate de M1-T5: consistencia del golden set frente al corpus, la ACL y la verdad del escenario.

    uv run app/evals/golden/validate_golden.py

Comprueba (sin BD ni LLM):
  1. Cada cita (doc_id, versión, locator) existe en el manifest y en el fichero:
     sección del DOCX (D1–D8 en 8D, título "Heading 1" en el resto), hoja y filas del XLSX, página del PDF.
  2. Cada `user` existe en acl.json y tiene ACL sobre todas las citas esperadas de su ítem
     (también las evidencias y filas del AMFE de los casos 8D).
  3. En las negativas de ACL, `forbidden_doc_ids` y `forbidden_folders` quedan fuera de la ACL del usuario.
  4. Recuentos por categoría (50 Q&A, 5 casos) y mínimos (OCR, AMFE sucio, casos de PLANT §9).
  5. Casos 8D: IDs de 8D en el manifest, reclamación en complaints/index.json, contención idéntica a
     scenario_truth.json, ≥2 hipótesis ligadas a fila del AMFE o marcadas fuera del AMFE, y todas las
     filas de `amfe_links` de la verdad cubiertas.
  6. Ningún texto del golden set contiene etiquetas de familia de recurrencia (PAT-004).
Sale con código 1 si hay algún error.
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter
from functools import lru_cache
from pathlib import Path

import openpyxl
from docx import Document
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
SYN = REPO / "app/data/synthetic"
DOCS = SYN / "docs"

MANIFEST = {d["doc_id"]: d for d in json.loads((DOCS / "manifest.json").read_text())["documents"]}
ACL = json.loads((SYN / "acl.json").read_text())
TRUTH = {c["complaint_id"]: c for c in json.loads((REPO / "app/evals/scenario_truth.json").read_text())["complaints"]}
COMPLAINTS = {c["complaint_id"]: c for c in json.loads((SYN / "complaints/index.json").read_text())["complaints"]}

QA_TOTAL, CASES_TOTAL = 50, 5
CATEGORY_RANGES = {  # (mín, máx)
    "factual": (28, 32),
    "recurrence": (4, 6),
    "multilingual": (4, 6),
    "acl_negative": (5, 5),
    "not_found": (5, 5),
}
MIN_OCR, MIN_DIRTY = 3, 3
FAMILY_RE = re.compile(r"\bfamil(?:ia|y|ies|ias)\s+(?:de\s+recurrencia\s+)?[AB]\b", re.IGNORECASE)
ROWS_RE = re.compile(r"^(\d+)-(\d+)$")

errors: list[str] = []


def err(where: str, msg: str) -> None:
    errors.append(f"{where}: {msg}")


# ---------------------------------------------------------------- lectura de ficheros (con caché)
@lru_cache(maxsize=None)
def docx_headings(path: Path) -> tuple[str, ...]:
    out = []
    for p in Document(path).paragraphs:
        style = (p.style.name if p.style is not None else "").lower()
        if style.startswith("heading 1") and p.text.strip():
            out.append(p.text.strip())
    return tuple(out)


@lru_cache(maxsize=None)
def xlsx_book(path: Path):
    return openpyxl.load_workbook(path, data_only=False)


def row_has_data(ws, r: int) -> bool:
    return any(c.value not in (None, "") for c in ws[r])


@lru_cache(maxsize=None)
def pdf_pages(path: Path) -> int:
    return len(PdfReader(path).pages)


def user_roles(user: str) -> set[str]:
    u = ACL["users"].get(user)
    return {u["role"]} if u else set()


def can_read(user: str, doc_id: str) -> bool:
    return bool(user_roles(user) & set(MANIFEST[doc_id]["acl_roles"]))


# ---------------------------------------------------------------- validación de una cita
def check_citation(where: str, cit: dict, user: str | None) -> None:
    doc_id, version, loc = cit.get("doc_id"), cit.get("version"), cit.get("locator")
    doc = MANIFEST.get(doc_id)
    if doc is None:
        return err(where, f"doc_id {doc_id!r} no está en el manifest")
    if version != doc["version"]:
        err(where, f"{doc_id}: versión {version!r} ≠ manifest {doc['version']!r}")
    if not isinstance(loc, dict) or not loc:
        return err(where, f"{doc_id}: locator vacío")
    path = DOCS / doc["path"]
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        if set(loc) != {"page"}:
            return err(where, f"{doc_id}: un PDF se cita con {{'page': n}}, no {loc}")
        n = pdf_pages(path)
        if not (isinstance(loc["page"], int) and 1 <= loc["page"] <= n):
            err(where, f"{doc_id}: página {loc['page']} fuera de 1..{n}")
    elif suffix == ".docx":
        if set(loc) != {"section"}:
            return err(where, f"{doc_id}: un DOCX se cita con {{'section': ...}}, no {loc}")
        sec = loc["section"]
        heads = docx_headings(path)
        if doc.get("sections"):  # 8D: D1–D8
            if sec not in doc["sections"]:
                err(where, f"{doc_id}: sección {sec!r} no está en el manifest {doc['sections']}")
            elif not any(h.startswith(sec) for h in heads):
                err(where, f"{doc_id}: no hay título 'Heading 1' que empiece por {sec!r}")
        elif sec not in heads:
            err(where, f"{doc_id}: sección {sec!r} no es un título del DOCX {list(heads)}")
    elif suffix == ".xlsx":
        if set(loc) != {"sheet", "rows"}:
            return err(where, f"{doc_id}: un XLSX se cita con {{'sheet', 'rows': 'a-b'}}, no {loc}")
        wb = xlsx_book(path)
        if loc["sheet"] not in wb.sheetnames:
            return err(where, f"{doc_id}: hoja {loc['sheet']!r} no existe {wb.sheetnames}")
        m = ROWS_RE.match(str(loc["rows"]))
        if not m:
            return err(where, f"{doc_id}: rows {loc['rows']!r} no tiene formato 'a-b'")
        a, b = int(m[1]), int(m[2])
        ws = wb[loc["sheet"]]
        if not (1 <= a <= b <= ws.max_row):
            return err(where, f"{doc_id}/{loc['sheet']}: filas {a}-{b} fuera de 1..{ws.max_row}")
        for r in {a, b}:
            if not row_has_data(ws, r):
                err(where, f"{doc_id}/{loc['sheet']}: la fila {r} está vacía")
    else:
        err(where, f"{doc_id}: extensión no soportada {suffix}")
    if user is not None and user in ACL["users"] and not can_read(user, doc_id):
        err(where, f"el usuario {user} no tiene ACL sobre {doc_id} ({doc['folder']})")


def check_amfe_link(where: str, link: dict, user: str) -> None:
    check_citation(where, {"doc_id": link.get("doc_id"), "version": link.get("version"),
                           "locator": {"sheet": link.get("sheet"), "rows": f"{link.get('row')}-{link.get('row')}"}}, user)
    if MANIFEST.get(link.get("doc_id"), {}).get("doc_type") != "AMFE":
        err(where, f"amfe_link apunta a {link.get('doc_id')}, que no es un AMFE")


def scan_family_labels(where: str, obj) -> None:
    text = json.dumps(obj, ensure_ascii=False)
    for m in FAMILY_RE.finditer(text):
        err(where, f"contiene una etiqueta de familia de recurrencia: {m.group(0)!r}")


# ---------------------------------------------------------------- qa.jsonl
def load_jsonl(path: Path) -> list[dict]:
    out = []
    for i, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if line.strip():
            try:
                out.append(json.loads(line))
            except json.JSONDecodeError as e:
                err(f"{path.name}:{i}", f"JSON inválido: {e}")
    return out


QA_REQUIRED = ["id", "user", "question", "expected_answer", "must_include", "expected_citations", "category", "difficulty"]


def validate_qa(items: list[dict]) -> Counter:
    cats: Counter = Counter()
    ids = [it.get("id") for it in items]
    for dup, n in Counter(ids).items():
        if n > 1:
            err("qa", f"id duplicado {dup}")
    ocr = dirty = 0
    acl_cases = set()
    for it in items:
        w = f"qa/{it.get('id')}"
        for k in QA_REQUIRED:
            if k not in it:
                err(w, f"falta el campo {k!r}")
        user, cat = it.get("user"), it.get("category")
        cats[cat] += 1
        if cat not in CATEGORY_RANGES:
            err(w, f"categoría desconocida {cat!r}")
        if it.get("difficulty") not in {"easy", "medium", "hard"}:
            err(w, f"difficulty inválida {it.get('difficulty')!r}")
        if user not in ACL["users"]:
            err(w, f"usuario {user!r} no existe en acl.json")
        cits = it.get("expected_citations", [])
        for j, cit in enumerate(cits):
            check_citation(f"{w}/cit{j}", cit, user)
        cited = {c["doc_id"] for c in cits}
        if cat in {"acl_negative", "not_found"}:
            if cits:
                err(w, f"{cat} debe tener expected_citations = []")
            if not re.search(r"no encontrado|not found|sin acceso", it.get("expected_answer", ""), re.I):
                err(w, "la respuesta esperada debe decir 'no encontrado' / 'sin acceso'")
        else:
            if not cits:
                err(w, "falta al menos una cita esperada")
            if not it.get("must_include"):
                err(w, "must_include vacío")
        if cat == "acl_negative":
            fb = it.get("forbidden_doc_ids") or []
            if not fb:
                err(w, "negativa de ACL sin forbidden_doc_ids")
            for d in fb:
                if d not in MANIFEST:
                    err(w, f"forbidden_doc_id {d} no está en el manifest")
                elif can_read(user, d):
                    err(w, f"forbidden_doc_id {d} SÍ es legible por {user}: la negativa no prueba nada")
            for f in it.get("forbidden_folders") or []:
                roles = ACL["folders"].get(f)
                if roles is None:
                    err(w, f"forbidden_folder {f} no existe en acl.json")
                elif user_roles(user) & set(roles):
                    err(w, f"forbidden_folder {f} SÍ es legible por {user}")
        for d in it.get("required_doc_ids") or []:
            if d not in cited:
                err(w, f"required_doc_id {d} no está entre las citas esperadas")
        if it.get("requires_ocr"):
            if any("scanned" in MANIFEST[c]["quality_flags"] for c in cited if "quality_flags" in MANIFEST.get(c, {})):
                ocr += 1
            else:
                err(w, "requires_ocr pero no cita un documento escaneado")
        if it.get("dirty_amfe"):
            if any("dirty" in MANIFEST[c].get("quality_flags", []) for c in cited if c in MANIFEST):
                dirty += 1
            else:
                err(w, "dirty_amfe pero no cita el AMFE sucio")
        if it.get("category") == "multilingual":
            if it.get("question_lang") not in {"es", "en"} or it.get("doc_lang") not in {"es", "en"} \
                    or it["question_lang"] == it["doc_lang"]:
                err(w, "multilingüe: question_lang y doc_lang deben existir y ser distintos")
            else:
                langs = {MANIFEST[c]["language"].lower() for c in cited if c in MANIFEST}
                if it["doc_lang"] not in langs:
                    err(w, f"doc_lang {it['doc_lang']} no coincide con el idioma de las fuentes {langs}")
        if "acl_case" in it:
            acl_cases.add(it["acl_case"])
        if re.search(r"famil", it.get("expected_answer", ""), re.I):
            err(w, "expected_answer menciona 'familia'")
        scan_family_labels(w, it)
    if len(items) != QA_TOTAL:
        err("qa", f"{len(items)} ítems (se esperan {QA_TOTAL})")
    for cat, (lo, hi) in CATEGORY_RANGES.items():
        if not lo <= cats[cat] <= hi:
            err("qa", f"categoría {cat}: {cats[cat]} ítems (rango {lo}–{hi})")
    if ocr < MIN_OCR:
        err("qa", f"solo {ocr} ítems con OCR (mín. {MIN_OCR})")
    if dirty < MIN_DIRTY:
        err("qa", f"solo {dirty} ítems sobre el AMFE sucio (mín. {MIN_DIRTY})")
    if acl_cases != {1, 2, 3, 4, 5}:
        err("qa", f"faltan casos de ACL de PLANT §9: {sorted({1, 2, 3, 4, 5} - acl_cases)}")
    cats["_ocr"], cats["_dirty_amfe"] = ocr, dirty
    return cats


# ---------------------------------------------------------------- 8d_cases.jsonl
CASE_REQUIRED = ["case_id", "complaint_file", "user", "expected_d1_team_roles", "expected_d2_facts",
                 "expected_d3_containment", "expected_similar_8d", "forbidden_similar_8d",
                 "expected_root_cause_hypotheses", "must_not", "max_seconds"]
CONTAINMENT_KEYS = ["lots", "shipments", "shipped_qty", "lots_in_stock", "lots_produced_after_notification",
                    "priority_lots_night_shift", "per_part", "minimum", "recommended", "verify_before_release",
                    "verify_next_nut_lots", "mt07_counter_at_end_of_complained_lot",
                    "over_scope_signal_ecoat_chemistry_lots"]


def is_8d(doc_id: str) -> bool:
    return MANIFEST.get(doc_id, {}).get("doc_type") == "8D"


def validate_cases(cases: list[dict]) -> None:
    seen = set()
    for cs in cases:
        w = f"8d/{cs.get('case_id')}"
        for k in CASE_REQUIRED:
            if k not in cs:
                err(w, f"falta el campo {k!r}")
        user = cs.get("user")
        if user != "inaki.calidad":
            err(w, f"user debe ser inaki.calidad, no {user!r}")
        if cs.get("max_seconds") != 180:
            err(w, "max_seconds debe ser 180")
        cid = cs.get("complaint_id")
        seen.add(cid)
        if cid not in COMPLAINTS:
            err(w, f"reclamación {cid} no está en complaints/index.json")
            continue
        fpath = REPO / cs.get("complaint_file", "")
        if not fpath.is_file() or fpath.name != COMPLAINTS[cid]["file"]:
            err(w, f"complaint_file {cs.get('complaint_file')} no existe o no coincide con el índice")
        truth = TRUTH.get(cid)
        if truth is None:
            err(w, f"{cid} no está en scenario_truth.json")
            continue
        # 8D similares
        for k in ("expected_similar_8d", "forbidden_similar_8d"):
            for d in cs.get(k, []):
                if not is_8d(d):
                    err(w, f"{k}: {d} no es un 8D del manifest")
        if cs.get("most_relevant_8d") and not is_8d(cs["most_relevant_8d"]):
            err(w, f"most_relevant_8d {cs['most_relevant_8d']} no es un 8D del manifest")
        if set(cs.get("expected_similar_8d", [])) != set(truth["recurrence"]["similar_8d_expected"]):
            err(w, "expected_similar_8d no coincide con scenario_truth")
        # contención
        d3 = cs.get("expected_d3_containment", {})
        ce = truth["containment_expected"]
        for k in CONTAINMENT_KEYS:
            if (k in ce) != (k in d3):
                err(w, f"contención: la clave {k!r} está en uno solo de golden/verdad")
            elif k in ce and ce[k] != d3[k]:
                err(w, f"contención: {k!r} no coincide con scenario_truth")
        # hipótesis
        hyps = cs.get("expected_root_cause_hypotheses", [])
        if len(hyps) < 2:
            err(w, f"solo {len(hyps)} hipótesis (mín. 2)")
        linked = set()
        for h in hyps:
            hw = f"{w}/{h.get('id')}"
            link = h.get("amfe_link")
            if link:
                check_amfe_link(hw, link, user)
                linked.add((link["doc_id"], link["sheet"], link["row"]))
                if h.get("outside_amfe"):
                    err(hw, "tiene amfe_link y outside_amfe = true a la vez")
            elif not h.get("outside_amfe"):
                err(hw, "sin amfe_link debe marcarse outside_amfe = true")
            for j, ev in enumerate(h.get("evidence", [])):
                check_citation(f"{hw}/ev{j}", ev, user)
        for al in truth["root_cause_expected"].get("amfe_links", []):
            if (al["doc_id"], al["sheet"], al["row"]) not in linked:
                err(w, f"la fila de verdad {al['doc_id']}/{al['sheet']}/{al['row']} no está en ninguna hipótesis")
        if not truth["root_cause_expected"].get("amfe_links") and not all(h.get("outside_amfe") for h in hyps):
            err(w, "la verdad no tiene AMFE: todas las hipótesis deben ir marcadas outside_amfe")
        if not cs.get("must_not"):
            err(w, "must_not vacío")
        scan_family_labels(w, cs)
    if len(cases) != CASES_TOTAL:
        err("8d", f"{len(cases)} casos (se esperan {CASES_TOTAL})")
    if seen != set(COMPLAINTS):
        err("8d", f"reclamaciones sin caso: {sorted(set(COMPLAINTS) - seen)}")


# ---------------------------------------------------------------- PAT-004: el golden set no entra al corpus
def validate_isolation() -> None:
    for d in MANIFEST.values():
        if "evals" in d["path"].split("/") or d["path"].endswith((".jsonl", ".py", ".md")):
            err("manifest", f"{d['doc_id']} apunta a {d['path']}: el golden set no puede ser ingerible")


def main() -> int:
    qa = load_jsonl(HERE / "qa.jsonl")
    cases = load_jsonl(HERE / "8d_cases.jsonl")
    cats = validate_qa(qa)
    validate_cases(cases)
    validate_isolation()
    print("qa.jsonl:", len(qa), "ítems ·", ", ".join(f"{k}={v}" for k, v in sorted(cats.items())))
    print("8d_cases.jsonl:", len(cases), "casos")
    n_cits = sum(len(i.get("expected_citations", [])) for i in qa)
    print("citas esperadas en qa:", n_cits)
    if errors:
        print(f"\nGATE M1-T5: FAIL ({len(errors)} errores)")
        for e in errors:
            print(" -", e)
        return 1
    print("\nGATE M1-T5: OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
