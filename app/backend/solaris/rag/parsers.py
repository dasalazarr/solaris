"""Parsing con chunking estructural (M2-T3, F01). Sin Docling (desviación de ADR-0002, ver traza).

Unidad de cita por formato (F02, M1-T2):
  * DOCX: sección por `Heading 1`. En los 8D, "D1"…"D8"; lo previo al primer encabezado es la
    cabecera (`header`). En IT/TPL/EVAL, el apartado ("1. Objeto").
  * PDF: página. Si la página no tiene capa de texto, OCR (rapidocr, ONNX) sobre la imagen embebida.
  * XLSX: hoja + bloque de filas (números de fila reales de la hoja).
Una unidad que supera `MAX_CHARS` se parte en trozos con el mismo locator y `part` = 1..n (e5 trunca
a 512 tokens). Todo el contenido extraído es DATO: nunca se interpreta (R03).
"""

import io
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

MAX_CHARS = 1800
XLSX_BLOCK_ROWS = 25
OCR_MIN_CHARS = 20  # por debajo, la página se trata como escaneada
_D_RE = re.compile(r"^\s*(D[1-8])\b")


@dataclass
class Chunk:
    content: str
    locator: dict[str, Any]
    page: int | None = None
    section: str | None = None


@dataclass
class Parsed:
    chunks: list[Chunk]
    pages: int | None = None
    ocr_pages: list[int] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


# --- utilidades ------------------------------------------------------------------------------


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _split(text: str, limit: int = MAX_CHARS) -> list[str]:
    """Parte por líneas (y, si una línea no cabe, por longitud) sin superar `limit`."""
    text = _clean(text)
    if len(text) <= limit:
        return [text] if text else []
    parts: list[str] = []
    cur = ""
    for line in text.split("\n"):
        while len(line) > limit:
            if cur:
                parts.append(cur)
                cur = ""
            parts.append(line[:limit])
            line = line[limit:]
        if cur and len(cur) + 1 + len(line) > limit:
            parts.append(cur)
            cur = line
        else:
            cur = f"{cur}\n{line}" if cur else line
    if cur.strip():
        parts.append(cur)
    return [p.strip() for p in parts if p.strip()]


def _emit(text: str, locator: dict[str, Any], page=None, section=None) -> list[Chunk]:
    pieces = _split(text)
    if len(pieces) == 1:
        return [Chunk(pieces[0], dict(locator), page, section)]
    return [
        Chunk(p, {**locator, "part": i}, page, section) for i, p in enumerate(pieces, start=1)
    ]


def cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:g}"
    return _clean(str(v))


# --- DOCX ------------------------------------------------------------------------------------


def _docx_blocks(path: Path) -> Iterator[tuple[str | None, str]]:
    """(heading_text | None, texto) en orden de documento; tablas como filas 'a | b | c'."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith("heading 1") or style.lower() == "title":
                yield text, ""
            else:
                prefix = "- " if "list" in style.lower() else ""
                yield None, prefix + text
        elif tag == "tbl":
            for row in Table(el, d).rows:
                cells: list[str] = []
                for c in row.cells:
                    t = _clean(c.text).replace("\n", " / ")
                    if t and (not cells or cells[-1] != t):  # celdas combinadas se repiten
                        cells.append(t)
                if cells:
                    yield None, " | ".join(cells)


def parse_docx(path: Path) -> Parsed:
    sections: list[tuple[str, str, list[str]]] = [("header", "", [])]  # (section, heading, lines)
    for heading, text in _docx_blocks(path):
        if heading is not None:
            m = _D_RE.match(heading)
            sections.append((m.group(1) if m else heading[:120], heading, [heading]))
        else:
            sections[-1][2].append(text)
    chunks: list[Chunk] = []
    for name, heading, lines in sections:
        body = "\n".join(lines)
        if not body.strip():
            continue
        loc: dict[str, Any] = {"section": name}
        if heading and heading != name:
            loc["heading"] = heading[:200]
        chunks.extend(_emit(body, loc, section=name))
    return Parsed(chunks)


# --- PDF + OCR -------------------------------------------------------------------------------


class OCRUnavailable(RuntimeError):
    pass


_ocr_engine = None


def _ocr_image(img) -> str:
    """OCR con rapidocr (ONNX, pip, sin binarios del sistema). Agrupa las cajas por línea."""
    global _ocr_engine
    try:
        import numpy as np
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise OCRUnavailable("rapidocr-onnxruntime no instalado") from exc
    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    result, _ = _ocr_engine(np.array(img.convert("RGB")))
    boxes = []
    for box, text, score in result or []:
        ys = [pt[1] for pt in box]
        xs = [pt[0] for pt in box]
        boxes.append((min(ys), max(ys), min(xs), str(text), float(score)))
    boxes.sort(key=lambda b: (b[0], b[2]))
    lines: list[list[tuple]] = []
    for b in boxes:
        if lines:
            last = lines[-1]
            top = min(x[0] for x in last)
            bottom = max(x[1] for x in last)
            mid = (b[0] + b[1]) / 2
            if top <= mid <= bottom:
                last.append(b)
                continue
        lines.append([b])
    return "\n".join(" ".join(x[3] for x in sorted(line, key=lambda x: x[2])) for line in lines)


def parse_pdf(path: Path, *, ocr: bool = True) -> Parsed:
    import pypdf
    from PIL import Image

    reader = pypdf.PdfReader(str(path))
    out = Parsed([], pages=len(reader.pages))
    for n, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        loc: dict[str, Any] = {"page": n}
        if len(text) < OCR_MIN_CHARS:
            images = list(page.images)
            if not images:
                out.notes.append(f"página {n}: sin texto ni imágenes")
                continue
            if not ocr:
                out.notes.append(f"página {n}: escaneada, OCR desactivado")
                continue
            img = max(
                (Image.open(io.BytesIO(im.data)) for im in images),
                key=lambda i: i.size[0] * i.size[1],
            )
            text = _clean(_ocr_image(img))
            loc["ocr"] = True
            out.ocr_pages.append(n)
            if not text:
                out.notes.append(f"página {n}: OCR sin texto")
                continue
        out.chunks.extend(_emit(text, loc, page=n))
    return out


# --- XLSX ------------------------------------------------------------------------------------


def load_grid(ws) -> dict[int, dict[int, Any]]:
    """{fila: {col: valor}} con los valores de las celdas combinadas copiados a todo el rango."""
    grid: dict[int, dict[int, Any]] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and cell_text(c.value):
                grid.setdefault(c.row, {})[c.column] = c.value
    for rng in ws.merged_cells.ranges:
        v = grid.get(rng.min_row, {}).get(rng.min_col)
        if v is None:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                grid.setdefault(r, {})[col] = v
    return grid


def _row_line(cells: dict[int, Any]) -> str:
    vals: list[str] = []
    for col in sorted(cells):
        t = cell_text(cells[col]).replace("\n", " / ")
        if t and (not vals or vals[-1] != t):
            vals.append(t)
    return " | ".join(vals)


def _header_row(grid: dict[int, dict[int, Any]]) -> int | None:
    best, best_n = None, 2
    for r in sorted(grid)[:20]:
        vals = list(grid[r].values())
        n = len({cell_text(v) for v in vals})
        if n > best_n and all(isinstance(v, str) for v in vals):
            best, best_n = r, n
    return best


def iter_sheets(path: Path) -> Iterable[tuple[Any, bool]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws, ws.sheet_state != "visible"
    finally:
        wb.close()


def parse_xlsx(path: Path) -> Parsed:
    out = Parsed([])
    for ws, hidden in iter_sheets(path):
        grid = load_grid(ws)
        if not grid:
            continue
        hdr = _header_row(grid)
        hdr_line = _row_line(grid[hdr]) if hdr else ""
        rows = sorted(grid)
        blocks: list[list[int]] = [[]]
        size = 0
        for r in rows:
            line_len = len(_row_line(grid[r])) + 12
            cur = blocks[-1]
            if cur and (len(cur) >= XLSX_BLOCK_ROWS or size + line_len > MAX_CHARS - 200):
                blocks.append([])
                size = 0
            blocks[-1].append(r)
            size += line_len
        for block in blocks:
            if not block:
                continue
            lines = [f"Hoja: {ws.title}"]
            if hdr and block[0] > hdr:
                lines.append(f"[cabecera fila {hdr}] {hdr_line}")
            lines += [f"[fila {r}] {_row_line(grid[r])}" for r in block]
            loc: dict[str, Any] = {"sheet": ws.title, "rows": [block[0], block[-1]]}
            if hidden:
                loc["hidden"] = True
            out.chunks.extend(_emit("\n".join(lines), loc, section=ws.title[:120]))
    return out


def parse(path: Path, *, ocr: bool = True) -> Parsed:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path, ocr=ocr)
    if suffix == ".xlsx":
        return parse_xlsx(path)
    raise ValueError(f"Formato no soportado: {path.suffix}")
