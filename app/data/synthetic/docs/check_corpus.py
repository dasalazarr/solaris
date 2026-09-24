# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "openpyxl>=3.1", "pypdf>=4.0"]
# ///
"""Gate de M1-T2: coherencia del corpus documental con el manifest, PLANT.md §11 y el ERP mock vivo.

Uso:  uv run app/data/synthetic/docs/check_corpus.py
Requiere el contenedor `solaris-db` levantado (docker compose up -d db). Sale con código 1 si algo falla.
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from docx import Document
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
PLANT = HERE.parent / "PLANT.md"
ACL = json.loads((HERE.parent / "acl.json").read_text(encoding="utf-8"))
DOC_EXT = {".docx", ".xlsx", ".pdf", ".eml", ".txt", ".md"}
TOP = ["calidad", "produccion", "compras", "direccion"]

FORBIDDEN = [
    (r"famil(ia|y)\s+[AB]\b", re.I), (r"\bfamilia de recurrencia\b", re.I), (r"\bEbro\b", re.I),
    (r"\bVW\b", 0), (r"\bVolkswagen\b", re.I), (r"\bSEAT\b", 0), (r"\bSeat\b", 0), (r"\bCupra\b", re.I), (r"\bAudi\b", re.I),
    (r"\b[SŠ]koda\b", re.I), (r"\bPorsche\b", re.I), (r"\bRenault\b", re.I), (r"\bStellantis\b", re.I), (r"\bPeugeot\b", re.I),
    (r"\bCitro[eë]n\b", re.I), (r"\bOpel\b", re.I), (r"\bMercedes\b", re.I), (r"\bBMW\b", 0), (r"\bFord\b", 0), (r"\bToyota\b", re.I),
    (r"\bTesla\b", re.I), (r"\bNissan\b", re.I), (r"\bHyundai\b", re.I), (r"\bVolvo\b", re.I), (r"\bGestamp\b", re.I),
    (r"\bBenteler\b", re.I), (r"\bMagna\b", 0), (r"\bFaurecia\b", re.I), (r"\bForvia\b", re.I), (r"\bLear\b", 0),
]

LOT_RE = re.compile(r"\bL\d{5}-AR\d{4}-\d{2}\b")
MAT_RE = re.compile(r"\bS-[A-Z]{4}-\d{6}\b")
SHIP_RE = re.compile(r"\bAL-\d{2}-\d{5}\b")
EIGHTD_RE = re.compile(r"\b8D-ARGA-\d{4}-\d{3}\b")

fails: list[str] = []


def check(cond: bool, msg: str) -> None:
    print(("  OK   " if cond else "  FAIL ") + msg)
    if not cond:
        fails.append(msg)


def sql(q: str) -> list[list[str]]:
    r = subprocess.run(["docker", "exec", "-i", "solaris-db", "psql", "-U", "solaris", "-d", "solaris", "-At", "-F", "|", "-c", q],
                       capture_output=True, text=True, timeout=60)
    if r.returncode:
        raise RuntimeError(r.stderr.strip())
    return [line.split("|") for line in r.stdout.strip().splitlines() if line]


def text_of(path: Path) -> str:
    if path.suffix == ".docx":
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts += [c.text for c in row.cells]
        for s in d.sections:
            parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
        return "\n".join(parts)
    if path.suffix == ".xlsx":
        wb = openpyxl.load_workbook(path)
        parts = []
        for ws in wb.worksheets:
            parts.append(ws.title)
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        parts.append(str(c.value))
                    if c.comment:
                        parts.append(c.comment.text)
        return "\n".join(parts)
    if path.suffix == ".pdf":
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    return path.read_text(encoding="utf-8", errors="replace")


def main() -> int:
    manifest = json.loads((HERE / "manifest.json").read_text(encoding="utf-8"))
    docs = manifest["documents"]

    print("1. Manifest ↔ disco")
    on_disk = {str(p.relative_to(HERE)) for t in TOP for p in (HERE / t).rglob("*") if p.is_file() and p.suffix in DOC_EXT}
    in_manifest = {d["path"] for d in docs}
    check(not (in_manifest - on_disk), f"todas las entradas del manifest existen en disco ({len(in_manifest)})"
          + (f" — faltan: {sorted(in_manifest - on_disk)}" if in_manifest - on_disk else ""))
    check(not (on_disk - in_manifest), "todo documento en disco está en el manifest"
          + (f" — sobran: {sorted(on_disk - in_manifest)}" if on_disk - in_manifest else ""))
    check(len({d["doc_id"] for d in docs}) == len(docs), "doc_id únicos")
    req = {"doc_id", "version", "title", "path", "folder", "doc_type", "language", "part_refs", "pages", "acl_roles"}
    check(all(req <= d.keys() for d in docs), "todas las entradas tienen los campos obligatorios")
    check(all(d["acl_roles"] == ACL["folders"][d["folder"]] for d in docs), "acl_roles = acl.json según la carpeta")
    check(all(d["path"].startswith(d["folder"] + "/") for d in docs), "path coherente con folder")

    texts = {d["path"]: text_of(HERE / d["path"]) for d in docs}

    print("2. 8D frente a PLANT.md §11")
    plant_ids = re.findall(r"^\| (8D-ARGA-\d{4}-\d{3}) \|", PLANT.read_text(encoding="utf-8"), re.M)
    eightd = [d for d in docs if d["doc_type"] == "8D"]
    check(len(plant_ids) == 15, f"PLANT.md §11 lista 15 8D ({len(plant_ids)})")
    check(sorted(d["doc_id"] for d in eightd) == sorted(plant_ids), f"exactamente 15 8D con los IDs de PLANT.md §11 ({len(eightd)})")
    langs = Counter(d["language"] for d in eightd)
    check(langs["EN"] > 0 and langs["ES"] > 0, f"8D en EN y ES ({dict(langs)})")

    print("3. Lotes y datos del ERP (base de datos viva)")
    try:
        lots = {r[0] for r in sql("select lot_code from erp.lots")}
        mats = {r[0] for r in sql("select lot_code from erp.material_lots")}
        ships = {r[0] for r in sql("select shipment_id from erp.shipments")}
        comp = {r[0]: r for r in sql("select report_8d_id, complaint_id, lot_code, part_ref, customer_code, received_date "
                                      "from erp.complaints where report_8d_id is not null")}
    except Exception as e:  # noqa: BLE001
        check(False, f"consulta al ERP vía docker ({e})")
        lots = mats = ships = set()
        comp = {}
    if comp:
        for d in eightd:
            t = texts[d["path"]]
            c = comp.get(d["doc_id"])
            if not c:
                check(False, f"{d['doc_id']}: sin reclamación en erp.complaints")
                continue
            _, cid, clot, cref, ccust, _ = c
            cited = set(LOT_RE.findall(t))
            bad = cited - lots
            ok = clot in cited and d.get("lot_code") == clot and cid in t and cref in d["part_refs"] and not bad
            check(ok, f"{d['doc_id']}: lote {clot} citado = erp.complaints; {len(cited)} lotes citados existen en erp.lots"
                  + (f" — inexistentes: {sorted(bad)}" if bad else ""))
        all_txt = "\n".join(texts.values())
        bad_m = set(MAT_RE.findall(all_txt)) - mats
        bad_s = set(SHIP_RE.findall(all_txt)) - ships
        bad_l = set(LOT_RE.findall(all_txt)) - lots
        check(not bad_m, f"lotes de material citados en todo el corpus existen en erp.material_lots ({len(set(MAT_RE.findall(all_txt)))})"
              + (f" — {sorted(bad_m)}" if bad_m else ""))
        check(not bad_s, f"albaranes citados existen en erp.shipments ({len(set(SHIP_RE.findall(all_txt)))})" + (f" — {sorted(bad_s)}" if bad_s else ""))
        check(not bad_l, f"lotes de producción citados en todo el corpus existen en erp.lots ({len(set(LOT_RE.findall(all_txt)))})"
              + (f" — {sorted(bad_l)}" if bad_l else ""))

    print("4. Términos prohibidos (familias, marcas reales)")
    hits = []
    for p, t in texts.items():
        for pat, fl in FORBIDDEN:
            for m in re.finditer(pat, t, fl):
                hits.append(f"{p}: '{m.group(0)}'")
    check(not hits, "ningún documento contiene 'familia A/B', 'Ebro' ni marcas reales de OEM/Tier 1" + (f" — {hits[:10]}" if hits else ""))

    print("5. Recuentos por tipo (DoD)")
    ct = Counter(d["doc_type"] for d in docs)
    check(ct["8D"] == 15, f"8D = 15 ({ct['8D']})")
    amfe = [d for d in docs if d["doc_type"] == "AMFE"]
    check(len(amfe) == 5 and all(d["path"].endswith(".xlsx") for d in amfe), f"AMFE = 5 en xlsx ({len(amfe)})")
    check(all(re.fullmatch(r"AMFE-[A-Z0-9-]+-\d{2}_v\d+\.xlsx", Path(d["path"]).name) for d in amfe), "nombres AMFE-<ref o área>-NN_vN.xlsx")
    check(sum("dirty" in d.get("quality_flags", []) for d in amfe) == 1, "1 AMFE 'sucio'")
    check({"AR-1003", "AR-1004"} <= set().union(*[set(d["part_refs"]) for d in amfe if d["doc_id"] == "AMFE-AR1003-01"]), "AMFE de AR-1003/1004 (CR-01)")
    t1003 = texts.get(next((d["path"] for d in amfe if d["doc_id"] == "AMFE-AR1003-01"), ""), "")
    check(all(k in t1003 for k in ("falta de fusión", "boquilla", "hilo", "Parámetros")), "AMFE AR-1003 cubre grieta/falta de fusión con causas boquilla, parámetros y lote de hilo")
    t1007 = texts.get(next((d["path"] for d in amfe if d["doc_id"] == "AMFE-AR1007-01"), ""), "")
    check("MT-07" in t1007 and "Rebaba" in t1007, "AMFE de AR-1007 (punzón MT-07, rebaba)")
    check(ct["PC"] == 5, f"planes de control = 5 ({ct['PC']})")
    its = [d for d in docs if d["doc_type"] == "IT"]
    check(len(its) == 8, f"instrucciones de trabajo = 8 ({len(its)})")
    scans = [d for d in its if "scanned" in d.get("quality_flags", [])]
    check(len(scans) == 1 and all(d["path"].endswith(".pdf") and not texts[d["path"]].strip() for d in scans),
          "1 IT en PDF escaneado sin capa de texto")
    cr01 = [d for d in its if d["doc_id"].startswith("IT-L2-CR01-")]
    check(len(cr01) == 1, "existe IT-L2-CR01-*")
    if cr01:
        t = texts[cr01[0]["path"]]
        check("8 h" in t and "mañana" in t and "tarde" in t and "22:00" not in t,
              "IT-L2-CR01 vigente: cambio de boquilla cada 8 h solo en mañana y tarde (sin 22:00 / noche)")
    mant = [d for d in docs if d["folder"] == "produccion/mantenimiento"]
    tm = "\n".join(texts[d["path"]] for d in mant)
    check(len(mant) >= 2 and "40.000" in tm and "60.000" in tm and "CR-01" in tm,
          "mantenimiento: plan de afilado (MT-07 40.000 → 60.000) + registro de boquillas de CR-01")
    evals = [d for d in docs if d["doc_type"] == "EVAL"]
    check(2 <= len(evals) <= 3 and any(d.get("supplier") == "S-GOIE" for d in evals), f"2–3 EVAL, una de S-GOIE ({len(evals)})")
    check(ct["COST"] >= 1 and all(d["folder"] == "direccion/confidencial" for d in docs if d["doc_type"] == "COST"), "COST-* en direccion/confidencial")
    tpl = {d["doc_id"] for d in docs if d["doc_type"] == "TPL"}
    check(tpl == {"TPL-8D-ARGA", "TPL-8D-OEMN"}, f"plantillas TPL-8D-ARGA y TPL-8D-OEMN ({sorted(tpl)})")

    print("6. Reglas del escenario (PAT-004)")
    refs_prev = {d["doc_id"]: set(EIGHTD_RE.findall(texts[d["path"]])) - {d["doc_id"]} for d in eightd}
    citing = sorted(k for k, v in refs_prev.items() if v)
    check(citing == ["8D-ARGA-2025-014"], f"solo 8D-ARGA-2025-014 cita un 8D anterior ({citing})")
    other = [d["path"] for d in docs if d["doc_type"] != "8D" and EIGHTD_RE.search(texts[d["path"]]) and d["folder"] != "direccion/confidencial"]
    check(not other, "ningún documento no-8D (salvo dirección) cita IDs de 8D" + (f" — {other}" if other else ""))

    print()
    if fails:
        print(f"GATE M1-T2: FAIL ({len(fails)} comprobaciones)")
        return 1
    print(f"GATE M1-T2: OK — {len(docs)} documentos, {dict(sorted(ct.items()))}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
