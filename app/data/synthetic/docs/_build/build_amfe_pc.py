"""AMFE de proceso (5, uno "sucio" — P06) y planes de control (5) en XLSX. DATOS SINTÉTICOS (ADR-0004)."""
from __future__ import annotations

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from common import BORDER, WRAP, new_workbook, save_wb, write_table, xlsx_header_block

AMFE_HEAD = ["Nº", "Paso de proceso", "Función / requisito", "Modo de fallo potencial", "Efecto potencial", "S", "Clase",
             "Causa potencial", "Control actual de prevención", "O", "Control actual de detección", "D", "NPR",
             "Acciones recomendadas", "Responsable / fecha", "Acciones tomadas", "S'", "O'", "D'", "NPR'"]
AMFE_W = [5, 20, 22, 24, 28, 4, 6, 30, 32, 4, 32, 4, 6, 28, 18, 28, 4, 4, 4, 6]


def _r(n, step, func, mode, eff, s, cls, cause, prev, o, det, d, rec="—", resp="—", taken="—", s2=None, o2=None, d2=None):
    npr = s * o * d
    npr2 = (s2 * o2 * d2) if s2 else None
    return [n, step, func, mode, eff, s, cls, cause, prev, o, det, d, npr, rec, resp, taken, s2, o2, d2, npr2]


# ----------------------------------------------------------------------------- AMFE AR-1003/1004 (CR-01)
AMFE_AR1003 = dict(
    code="AMFE-AR1003-01", version="v4", title="AMFE de proceso — Soporte de la bandeja de baterías LH/RH",
    refs=["AR-1003", "AR-1004"],
    header=[("AMFE de proceso", "AMFE-AR1003-01  ·  versión v4"), ("Piezas", "AR-1003 (LH), AR-1004 (RH) — cliente OEM Norte"),
            ("Líneas / equipos", "L1 PR-400 (MT-03, MT-04) → L2 CR-01 (MIG robot) → L3 EC-01 / VI-01"),
            ("Fecha original", "15/06/2023 (PPAP)"), ("Última revisión", "18/12/2025"),
            ("Equipo AMFE", "O. Zabaleta (Ingeniería L2), M. Irigoyen (Calidad), X. Urrutia (Producción), K. Arbizu (Mantenimiento), U. Ezkurra (Matricería)"),
            ("Escalas", "S, O, D 1–10 (manual AMFE interno PROC-CAL-05). NPR = S × O × D. Revisar si S ≥ 9 o NPR ≥ 100")],
    rows=[
        _r(10, "Recepción de bobina HSLA 420 (S-ULTZ)", "Material según especificación", "Propiedades mecánicas fuera de especificación",
           "Soldabilidad y resistencia reducidas", 8, "", "Error del proveedor en la colada", "Certificado 3.1 por bobina", 2,
           "Revisión del certificado en recepción", 5),
        _r(20, "Estampación transfer PR-400 (MT-03 / MT-04)", "Forma y cotas según plano", "Holgura de encaje con el refuerzo > 1,0 mm",
           "Falta de fusión por separación de chapas en W2", 8, "", "Desgaste de la matriz / radio de doblado", "Mantenimiento preventivo de matriz", 3,
           "Control en útil de control 1 pieza/turno", 4),
        _r(20, "Estampación transfer PR-400 (MT-03 / MT-04)", "Forma y cotas según plano", "Grieta de estampación en radio",
           "Rotura en servicio", 9, "", "Radio de matriz dañado", "Inspección de matriz en cada montaje", 2,
           "Visual 100 % en fin de línea L1", 5),
        _r(30, "Soldadura MIG robotizada CR-01 (cordones W1, W2)", "Resistencia de la soldadura ≥ 18 kN (arrancamiento)",
           "Grieta / falta de fusión en el cordón W2", "Rotura del soporte de la bandeja de baterías; riesgo para la sujeción de la batería (seguridad)",
           10, "CC", "Parámetros de soldadura fuera de la ventana del WPS (modificación en consola)",
           "Parámetros bloqueados con contraseña; verificación al inicio de turno (IT-L2-CR01-03)", 2,
           "Arrancamiento 3 piezas/lote (≥ 18 kN, límite de aviso 20 kN) + macrografía semanal", 4,
           "Mantener", "—", "Bloqueo de parámetros (06/2024)"),
        _r(30, "Soldadura MIG robotizada CR-01 (cordones W1, W2)", "Resistencia de la soldadura ≥ 18 kN (arrancamiento)",
           "Grieta / falta de fusión en el cordón W2", "Rotura del soporte de la bandeja de baterías; riesgo para la sujeción de la batería (seguridad)",
           10, "CC", "Desgaste de la boquilla / tubo de contacto (arco inestable, mala protección gaseosa)",
           "Cambio de boquilla y tubo de contacto cada 8 h (IT-L2-CR01-03); registro REG-L2-CR01-01", 3,
           "Arrancamiento 3 piezas/lote + macrografía semanal; límite de aviso 20 kN", 4,
           "Estudio de monitorización de estabilidad del arco", "O. Zabaleta / 06/2026", "Cambio cada 8 h implantado 03/2025"),
        _r(30, "Soldadura MIG robotizada CR-01 (cordones W1, W2)", "Resistencia de la soldadura ≥ 18 kN (arrancamiento)",
           "Grieta / falta de fusión en el cordón W2", "Rotura del soporte de la bandeja de baterías; riesgo para la sujeción de la batería (seguridad)",
           10, "CC", "Cambio de lote de hilo MIG (S-GOIE) con distinta respuesta del arco",
           "Hilo G3Si1 Ø1,0 homologado; certificado 3.1 por lote", 2, "Arrancamiento 3 piezas/lote + macrografía semanal", 4),
        _r(30, "Soldadura MIG robotizada CR-01 (cordones W1, W2)", "Resistencia de la soldadura ≥ 18 kN (arrancamiento)",
           "Grieta / falta de fusión en el cordón W2", "Rotura del soporte de la bandeja de baterías; riesgo para la sujeción de la batería (seguridad)",
           10, "CC", "Caudal de gas de protección insuficiente", "Caudalímetro con alarma a 12 l/min", 2,
           "Verificación de caudal al inicio de turno", 3),
        _r(30, "Soldadura MIG robotizada CR-01 (cordones W1, W2)", "Posición del cordón según plano", "Cordón desplazado > 2 mm",
           "Resistencia reducida; interferencia en montaje", 7, "", "Pieza mal posicionada en el útil", "Detectores de presencia y bridas neumáticas", 2,
           "Visual 100 % por el operario", 5),
        _r(30, "Soldadura MIG robotizada CR-01 (cordones W1, W2)", "Superficie de apoyo libre de salpicaduras", "Salpicaduras en la zona de apoyo",
           "Mal asiento en la bandeja; ruido", 4, "", "Boquilla sucia; antiproyecciones insuficiente", "Escariador de boquilla cada 20 piezas", 4,
           "Visual 100 %", 5),
        _r(40, "Control de soldadura (laboratorio)", "Detección de soldaduras NOK", "Lote NOK liberado",
           "Pieza NOK en cliente", 10, "CC", "Muestreo insuficiente para defectos intermitentes", "Plan de control PC-AR1003-01", 3,
           "Arrancamiento 3 piezas/lote; límite de aviso 20 kN → 5 piezas + macro", 5),
        _r(50, "Cataforesis EC-01 + curado HC-01", "Espesor 18–25 µm", "Espesor insuficiente en el cordón", "Corrosión prematura",
           6, "", "Efecto jaula de Faraday en el ángulo del refuerzo", "Posición de cuelgue definida", 3, "Medición de espesor 5 puntos, 1/20", 4),
        _r(60, "Inspección final VI-01 y embalaje", "Identificación LH/RH", "Mezcla de AR-1003 y AR-1004",
           "Parada de línea en cliente", 8, "", "Piezas simétricas en la misma zona", "Contenedores con color por lado", 2,
           "Visión VI-01 (forma)", 2),
    ],
    history=[("v1", "15/06/2023", "Emisión para PPAP", "O. Zabaleta"),
             ("v2", "14/06/2024", "Nueva causa: modificación de parámetros en consola (reclamación de cliente por grieta en W2). O 5→3, D 6→4 con arrancamiento por lote", "O. Zabaleta"),
             ("v3", "07/03/2025", "Nueva causa: desgaste de boquilla / tubo de contacto (reclamación de cliente por falta de fusión). Prevención: cambio cada 8 h. O 6→3", "O. Zabaleta"),
             ("v4", "18/12/2025", "Revisión de la detección: límite de aviso de 20 kN en arrancamiento", "O. Zabaleta")],
)

# ----------------------------------------------------------------------------- AMFE AR-1007 (MT-07)
AMFE_AR1007 = dict(
    code="AMFE-AR1007-01", version="v3", title="AMFE de proceso — Brida de anclaje del cinturón de seguridad",
    refs=["AR-1007"],
    header=[("AMFE de proceso", "AMFE-AR1007-01  ·  versión v3"), ("Pieza", "AR-1007 — cliente Leitzaran Seating"),
            ("Líneas / equipos", "L1 PR-250, matriz progresiva MT-07 (estaciones P1–P6) → L3 EC-01 / VI-01"),
            ("Fecha original", "22/05/2023 (PPAP)"), ("Última revisión", "30/09/2025"),
            ("Equipo AMFE", "U. Ezkurra (Matricería), M. Irigoyen (Calidad), J. Elizalde (Producción), K. Arbizu (Mantenimiento)"),
            ("Escalas", "S, O, D 1–10 (PROC-CAL-05). NPR = S × O × D")],
    rows=[
        _r(10, "Recepción de bobina HSLA 500 (S-ULTZ)", "Rm 560–700 MPa (acuerdo AR-1007: Rm ≤ 650 MPa)", "Material con Rm en el límite superior",
           "Desgaste acelerado del punzón; rebaba", 7, "", "Variación de colada del proveedor", "Acuerdo técnico con S-ULTZ (Rm ≤ 650 MPa)", 2,
           "Revisión del certificado 3.1", 5),
        _r(20, "Punzonado del agujero de anclaje Ø11,20 (MT-07, estación P3)", "Rebaba ≤ 0,2 mm", "Rebaba > 0,2 mm en el agujero de anclaje",
           "El tornillo de anclaje del cinturón no asienta; riesgo en el anclaje (seguridad)", 9, "CC",
           "Desgaste del punzón P3",
           "Afilado preventivo cada 60.000 golpes por S-ZIDA (PROC-L1-AFILADO-01 v3); alarma de contador en PR-250; punzón de repuesto", 3,
           "Galga de rebaba y medición con reloj en 4 posiciones cada 2 h", 4,
           "Ensayo de punzón recubierto (PM + TiCN)", "U. Ezkurra / 12/2025", "Pedido realizado; entrega retrasada"),
        _r(20, "Punzonado del agujero de anclaje Ø11,20 (MT-07, estación P3)", "Ø11,20 +0,15/0", "Ø del agujero por debajo de la tolerancia",
           "El tornillo de anclaje no entra; parada de línea en cliente", 8, "CC", "Desgaste y microastillado del punzón P3",
           "Afilado preventivo cada 60.000 golpes (PROC-L1-AFILADO-01 v3)", 3, "Calibre tampón pasa/no pasa Ø11,20 cada 2 h", 4),
        _r(20, "Punzonado del agujero de anclaje Ø11,20 (MT-07, estación P3)", "Ø11,20 +0,15/0", "Ø del agujero por debajo de la tolerancia",
           "El tornillo de anclaje no entra; parada de línea en cliente", 8, "CC", "Calibre tampón desgastado acepta agujeros pequeños",
           "Calibración cada 3 meses; control de desgaste semanal", 2, "—", 6),
        _r(20, "Punzonado del agujero de anclaje Ø11,20 (MT-07, estación P3)", "Rebaba ≤ 0,2 mm", "Rebaba > 0,2 mm en el agujero de anclaje",
           "El tornillo de anclaje del cinturón no asienta", 9, "CC", "Holgura punzón-matriz incorrecta tras el afilado",
           "Verificación de holgura a la recepción del afilado", 2, "Primera pieza tras montaje: galga y tampón", 3),
        _r(30, "Doblado de alas (MT-07, P5)", "Ángulo 90 ± 0,5°", "Ángulo fuera de tolerancia", "Montaje forzado en el asiento", 6, "",
           "Recuperación elástica variable", "Estación de reestampado", 3, "Útil de control 1 pieza/turno", 4),
        _r(40, "Cataforesis e inspección final", "Agujero libre de pintura acumulada", "Acumulación de e-coat en el agujero",
           "Tornillo con apriete forzado", 5, "", "Gota en el borde del agujero", "Posición de cuelgue", 3, "Tampón en inspección final 1/50", 5),
    ],
    history=[("v1", "22/05/2023", "Emisión para PPAP. Afilado de MT-07 cada 40.000 golpes", "U. Ezkurra"),
             ("v2", "22/11/2024", "Causa 'afilado fuera de plazo': alarma de contador y punzón de repuesto (reclamación de cliente por rebaba). Control cada 2 h", "U. Ezkurra"),
             ("v3", "30/09/2025", "Control de prevención actualizado al intervalo vigente del plan de afilado (60.000 golpes). Nuevas causas: material en el límite superior de Rm y desgaste del calibre (reclamación de cliente por Ø bajo)", "U. Ezkurra")],
)

# ----------------------------------------------------------------------------- AMFE AR-1002/AR-1005 (PR-400)
AMFE_AR1002 = dict(
    code="AMFE-AR1002-01", version="v4", title="AMFE de proceso — Estampación PR-400 (escuadra del parachoques y refuerzo del travesaño)",
    refs=["AR-1002", "AR-1005"],
    header=[("AMFE de proceso", "AMFE-AR1002-01  ·  versión v4"), ("Piezas", "AR-1002 y AR-1005 — cliente Ribera Chassis Systems"),
            ("Líneas / equipos", "L1 PR-400 (MT-02, MT-05) → L2 CR-02 (solo AR-1005) → L3 EC-01 / VI-01 / embalaje"),
            ("Fecha original", "10/03/2023"), ("Última revisión", "20/06/2025"),
            ("Equipo AMFE", "M. Irigoyen (Calidad), U. Ezkurra (Matricería), A. Olaberri (Logística), L. Sarasola (Compras)"),
            ("Escalas", "S, O, D 1–10 (PROC-CAL-05)")],
    rows=[
        _r(10, "Recepción de bobina (DC04 / HSLA 340)", "Re según especificación; AR-1005: Re ≥ 360 MPa", "Límite elástico en el mínimo de especificación",
           "Deformación permanente del refuerzo bajo carga (AR-1005)", 8, "CC", "Variación de colada", "Acuerdo técnico con S-ULTZ (ventana restringida AR-1005)", 2,
           "Ensayo de tracción en recepción por bobina", 3),
        _r(20, "Estampación MT-02 (AR-1002)", "Planitud ≤ 0,5 mm", "Planitud fuera de tolerancia", "La escuadra no apoya; montaje forzado", 6, "SC",
           "Inserto de calibrado desgastado o fuera de cota tras reparación", "Liberación de matriz tras reparación con medición de primera pieza", 2,
           "Planitud en mármol 1 pieza cada 2 h", 4),
        _r(20, "Estampación MT-05 (AR-1005)", "Espesor y radios según plano", "Adelgazamiento en el radio > 15 %", "Resistencia reducida", 8, "CC",
           "Radio de matriz dañado", "Inspección de matriz en cada montaje", 2, "Medición de espesor 1 pieza/lote", 5),
        _r(30, "Soldadura MIG CR-02 (AR-1005)", "Penetración ≥ 1,0 mm", "Falta de penetración", "Rotura del refuerzo", 9, "CC",
           "Parámetros fuera de ventana", "Parámetros bloqueados; verificación al inicio de turno", 2, "Macrografía semanal", 4),
        _r(50, "Embalaje y expedición", "Pieza sin deformación en recepción del cliente", "Deformación en transporte",
           "Rechazo en recepción", 6, "", "Separador de cartón inadecuado o apilado excesivo",
           "Separador de 5 mm validado; máx. 4 alturas; cambios de S-ERRO bajo gestión de cambios", 2, "Visual del embalaje en expedición", 4),
    ],
    history=[("v1", "10/03/2023", "Emisión", "M. Irigoyen"),
             ("v2", "25/07/2024", "Nuevo modo 'deformación en transporte' (reclamación de cliente)", "M. Irigoyen"),
             ("v3", "15/05/2025", "Causa 'material en el mínimo de especificación' (AR-1005)", "M. Irigoyen"),
             ("v4", "20/06/2025", "Causa 'herramienta reparada sin liberación dimensional' (AR-1002)", "M. Irigoyen")],
)

# ----------------------------------------------------------------------------- AMFE e-coat (L3)
AMFE_EC01 = dict(
    code="AMFE-L3-EC01-01", version="v2", title="AMFE de proceso — Cataforesis EC-01 y curado HC-01",
    refs=["AR-1001", "AR-1002", "AR-1003", "AR-1004", "AR-1005", "AR-1006", "AR-1007", "AR-1008", "AR-1009", "AR-1010", "AR-1011", "AR-1012"],
    header=[("AMFE de proceso", "AMFE-L3-EC01-01  ·  versión v2"), ("Alcance", "Todas las referencias (L3)"),
            ("Equipos", "Pretratamiento, EC-01, horno HC-01"), ("Fecha original", "02/02/2023"), ("Última revisión", "26/08/2024"),
            ("Equipo AMFE", "G. Lizarraga (Ingeniería L3), M. Irigoyen (Calidad), L. Sarasola (Compras)"), ("Escalas", "S, O, D 1–10")],
    rows=[
        _r(10, "Desengrase y fosfatado", "Superficie limpia", "Contaminación superficial", "Poros y falta de adherencia", 7, "",
           "Baño agotado", "Titulación por turno", 3, "Ensayo de rotura de gota de agua 1/turno", 5),
        _r(20, "Aclarado con agua desionizada", "Conductividad < 20 µS/cm", "Arrastre de sales al baño de e-coat", "Cráteres y poros", 7, "SC",
           "Conductividad alta", "Alarma a 20 µS/cm con parada", 2, "Registro continuo", 3),
        _r(30, "Cataforesis EC-01", "Espesor 18–25 µm", "Poros y cráteres en el e-coat", "Corrosión (niebla salina < 480 h)", 7, "SC",
           "Contaminación del baño por química de reposición (S-ARAK)", "Control de fosfatos y conductividad de cada lote de química antes de añadir", 2,
           "Inspección visual 1.000 lux 5 piezas/bastidor + niebla salina mensual", 4),
        _r(30, "Cataforesis EC-01", "Espesor 18–25 µm", "Espesor insuficiente", "Corrosión prematura", 7, "SC",
           "Tensión / tiempo de deposición bajos", "Receta por referencia bloqueada", 2, "Espesor 5 puntos, 1/20", 4),
        _r(40, "Curado HC-01", "180 °C / 20 min", "Curado insuficiente", "Falta de adherencia", 7, "",
           "Perfil de horno fuera de ventana", "Registro de temperatura continuo", 2, "Ensayo de MEK 1/turno", 4),
    ],
    history=[("v1", "02/02/2023", "Emisión", "G. Lizarraga"),
             ("v2", "26/08/2024", "Causas 'contaminación por química de reposición' y 'conductividad de aclarado' (reclamación de cliente por poros)", "G. Lizarraga")],
)


def _amfe_clean(spec) -> dict:
    wb = new_workbook()
    ws = wb.active
    ws.title = "AMFE"
    ws.cell(row=1, column=1, value=spec["title"]).font = Font(bold=True, size=13, color="1F4E79")
    top = xlsx_header_block(ws, spec["header"], top=2)
    write_table(ws, top, AMFE_HEAD, spec["rows"], AMFE_W)
    ws.freeze_panes = ws.cell(row=top + 1, column=4)
    ws2 = wb.create_sheet("Historial de revisiones")
    write_table(ws2, 1, ["Versión", "Fecha", "Cambio", "Autor"], spec["history"], [8, 12, 90, 16])
    rel = f"calidad/amfe/{spec['code']}_{spec['version']}.xlsx"
    save_wb(wb, rel)
    return {"doc_id": spec["code"], "version": spec["version"], "title": spec["title"], "path": rel,
            "folder": "calidad/amfe", "doc_type": "AMFE", "language": "ES", "part_refs": spec["refs"], "pages": None,
            "sheets": ["AMFE", "Historial de revisiones"]}


# ----------------------------------------------------------------------------- AMFE "sucio" AR-1006/AR-1010 (P06)
def _amfe_dirty() -> dict:
    wb = new_workbook(creator="oihana.z")
    ws = wb.active
    ws.title = "AMFE tuercas"
    yellow = PatternFill("solid", fgColor="FFFF99")
    ws.merge_cells("B2:P2")
    ws["B2"] = "AMFE PROCESO - SOLDADURA TUERCAS (AR-1006 / AR-1010)   rev 3 ??  (ult. mod. julio-25)"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B3"] = "Equipo: Oihana, Maialen, Koldo (mant.), Xabi"
    ws["H3"] = "Fecha orig: 3-mar-23"
    ws["K3"] = "Cliente: OEM Norte"
    ws["B4"] = "NOTA: las columnas S/O/D de la v2 están en la hoja 'viejo' - NO BORRAR"
    ws["B4"].font = Font(italic=True, color="C00000")
    # cabecera en dos filas con celdas combinadas, desplazada a la columna B y fila 6
    ws.merge_cells("B6:B7"); ws["B6"] = "Op."
    ws.merge_cells("C6:C7"); ws["C6"] = "Paso / función"
    ws.merge_cells("D6:D7"); ws["D6"] = "Modo de fallo"
    ws.merge_cells("E6:E7"); ws["E6"] = "Efecto(s)"
    ws.merge_cells("F6:F7"); ws["F6"] = "Sev"
    ws.merge_cells("G6:G7"); ws["G6"] = "Causa(s)"
    ws.merge_cells("H6:J6"); ws["H6"] = "Situación actual"
    ws["H7"] = "Prevención"; ws["I7"] = "Ocurr."; ws["J7"] = "Detección"
    ws.merge_cells("K6:K7"); ws["K6"] = "Det"
    ws.merge_cells("L6:L7"); ws["L6"] = "NPR"
    ws.merge_cells("M6:M7"); ws["M6"] = "Acción recomendada / resp."
    ws.merge_cells("N6:P6"); ws["N6"] = "Tras acciones"
    ws["N7"] = "S"; ws["O7"] = "O"; ws["P7"] = "D"
    for row in ws.iter_rows(min_row=6, max_row=7, min_col=2, max_col=16):
        for c in row:
            c.fill = PatternFill("solid", fgColor="D9D9D9")
            c.font = Font(bold=True, size=9)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, horizontal="center")
    data = [
        # fila, op, paso, modo, efecto, S, causa, prev, O, det, D, NPR, acción, S', O', D'
        (8, "10", "Alimentación tuerca M8 (AT-01)", "Tuerca no presente / invertida", "Tuerca faltante en cliente -> parada línea", "7",
         "atasco alimentador", "sensor presencia tuerca en electrodo", "2", "sensor + visual", "3", "=F8*I8*K8", "-", None, None, None),
        (9, None, None, "Tuerca de otra métrica (M6)", "no monta", 7, "mezcla en tolva", "tolva dedicada M8 (etiqueta)", "3*",
         "visual", "5 (ver nota)", "105", "poka-yoke diámetro?? pdte Koldo", None, None, None),
        (11, "20", "Soldadura por proyección tuerca\n(SP-01 AR-1006 / estación tuercas CR-02 AR-1010)", "Par de arrancamiento bajo / tuerca gira",
         "Tuerca gira al apretar compresor o inversor. Reclamación cliente", "8 (ver nota)",
         "Electrodo desalineado tras mantenimiento", "Galga alineación + prueba par 1ª pieza tras mant. (desde dic-24)", "2",
         "par arrancamiento 3 pzs / 2h", "4", "64", "ok", 8, 2, 4),
        (12, None, None, None, None, "8",
         "Recubrimiento Zn-Ni de la tuerca en límite sup. (S-BIDA)", "XRF en recepción por lote; ventana 8-11 µm acordada con BIDA", "2-3",
         "par arrancamiento 3 pzs / 2h", " 4", "~80", "revisar tras 8D tuercas inversor.  OJO: el lote de agosto-26 no se ha medido todavía (XRF averiado)", None, None, None),
        (13, None, None, None, None, "alta",
         "Corriente / tiempo de soldadura bajos", "programa bloqueado", "3", "monitorización corriente (solo SP-01!!)", "5", "#¡VALOR!", "", None, None, None),
        (15, "20", "Soldadura por proyección tuerca", "Rosca dañada / salpicadura en rosca", "El tornillo no entra", "6",
         "proyección mal fundida, expulsión", "parámetros", "4", "tampón roscado 1/turno", "4", "96", "Aumentar frec. tampón a cada 2h (M.I.)", 6, 3, 3),
        (17, "30", "Control par (lab)", "lote NOK liberado", "cliente", "8", "muestra pequeña", "PC-AR1006-01 v3", "3", "3 piezas cada 2h", "5",
         "=F17*I17*K17", "", None, None, None),
        (19, None, "e-coat (ver AMFE-L3-EC01)", None, None, None, None, None, None, None, None, None, None, None, None, None),
    ]
    for r in data:
        row, vals = r[0], r[1:]
        for j, v in enumerate(vals):
            if v is None:
                continue
            c = ws.cell(row=row, column=2 + j, value=v)
            c.alignment = WRAP
            c.font = Font(size=9)
            c.border = BORDER
    ws.merge_cells("B11:B13"); ws.merge_cells("C11:C13"); ws.merge_cells("D11:D13"); ws.merge_cells("E11:E13")
    ws.merge_cells("B8:B9"); ws.merge_cells("C8:C9")
    for addr in ("F11", "I12", "K9", "L12", "F13"):
        ws[addr].fill = yellow
    ws["F11"].comment = Comment("S=8 porque es SC, no CC (confirmado con cliente 2023). Maialen", "maialen.i")
    ws["L13"].comment = Comment("fórmula rota al pegar desde la v2", "oihana.z")
    ws["G21"] = "Pendiente: añadir AR-1010 a todas las filas (antes solo AR-1006)"
    ws["G22"] = "última revisión con equipo: 22/07/2025"
    ws["M24"] = "firmado: ___________"
    for col, w in zip("ABCDEFGHIJKLMNOP", [2, 6, 22, 22, 26, 8, 30, 32, 8, 26, 8, 8, 34, 5, 5, 5]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["Q"].hidden = True
    ws["Q11"] = "S=7 O=4 D=6 (v1)"
    # hoja "viejo" con restos de la v2
    ws2 = wb.create_sheet("viejo")
    ws2["A1"] = "copia v2 (dic-2024) - NO USAR"
    for i, r in enumerate([("Op", "Causa", "S", "O", "D", "NPR"), ("20", "electrodo desalineado", 8, 5, 6, 240),
                           ("20", "recubrimiento tuerca", 8, 2, 6, 96), ("", "", "", "", "", "")], start=3):
        for j, v in enumerate(r, start=1):
            ws2.cell(row=i, column=j, value=v)
    ws3 = wb.create_sheet("Hoja2")
    ws3["C5"] = "par mín 40 Nm   (plano AR-1006 rev C)  / AR-1010 idem?"
    ws3["C7"] = "XRF: 9,5 - 10,1 - 11,8 - 12,0 µm"
    rel = "calidad/amfe/AMFE-AR1006-01_v3.xlsx"
    save_wb(wb, rel)
    return {"doc_id": "AMFE-AR1006-01", "version": "v3", "title": "AMFE de proceso — Soldadura de tuercas por proyección (AR-1006 / AR-1010)",
            "path": rel, "folder": "calidad/amfe", "doc_type": "AMFE", "language": "ES", "part_refs": ["AR-1006", "AR-1010"],
            "pages": None, "sheets": ["AMFE tuercas", "viejo", "Hoja2"], "quality_flags": ["dirty"]}


# ----------------------------------------------------------------------------- Planes de control
PC_HEAD = ["Nº op.", "Proceso / operación", "Máquina / útil", "Nº", "Característica de producto", "Característica de proceso", "Clase",
           "Especificación / tolerancia", "Técnica de evaluación", "Muestra", "Frecuencia", "Método de control / registro", "Plan de reacción"]
PC_W = [6, 22, 16, 5, 24, 24, 6, 24, 22, 8, 22, 26, 34]

PCS = [
    dict(code="PC-AR1003-01", version="v4", refs=["AR-1003", "AR-1004"], title="Plan de control — Soporte de la bandeja de baterías LH/RH",
         header=[("Plan de control", "PC-AR1003-01 · v4 · Serie"), ("Piezas", "AR-1003, AR-1004 (OEM Norte)"), ("AMFE asociado", "AMFE-AR1003-01 v4"),
                 ("Fecha", "18/12/2025"), ("Aprobado", "I. Etxeberria (Calidad) · aprobación del cliente: OEM Norte SQE 05/01/2026")],
         rows=[
             [10, "Recepción de bobina", "—", 1, "Propiedades mecánicas", "", "", "HSLA 420 según EN 10149-2", "Certificado 3.1", "1", "Cada bobina", "Registro de recepción", "Bloquear bobina; aviso a Compras"],
             [20, "Estampación transfer", "PR-400 / MT-03, MT-04", 2, "Holgura de encaje con el refuerzo", "", "", "≤ 1,0 mm", "Útil de control CF-AR1003", "1", "Inicio de turno y cada 4 h", "Hoja de control L1", "Parar; ajustar matriz; segregar desde último OK"],
             [30, "Soldadura MIG", "CR-01", 3, "", "Parámetros WPS-CR01-03 (WFS 9,4 ± 0,3 m/min; 23,5 ± 0,5 V)", "", "Ventana WPS", "Lectura en consola (bloqueada)", "—", "Inicio de turno", "Checklist de célula (IT-L2-CR01-03)", "No arrancar; avisar a ingeniería L2"],
             [30, "Soldadura MIG", "CR-01", 4, "", "Caudal de gas M21", "", "14–18 l/min", "Caudalímetro en boquilla", "—", "Inicio de turno", "Checklist de célula", "Corregir; revisar piezas desde el último control"],
             [30, "Soldadura MIG", "CR-01", 5, "", "Estado de boquilla y tubo de contacto", "", "Cambio preventivo", "Visual + cambio", "—", "Cada 8 h (según IT-L2-CR01-03)", "Registro REG-L2-CR01-01", "Cambiar; si hay salpicadura excesiva, avisar al jefe de turno"],
             [30, "Soldadura MIG", "CR-01", 6, "", "Lote de hilo MIG en uso", "", "G3Si1 Ø1,0 homologado", "Registro de lote", "—", "En cada cambio de bobina", "Registro REG-L2-CR01-01", "Bloquear bobina no homologada"],
             [30, "Soldadura MIG", "CR-01", 7, "Aspecto del cordón W1/W2", "", "", "Sin grietas, poros ni mordeduras", "Visual", "100 %", "Continuo", "Autocontrol del operario", "Segregar y avisar"],
             [40, "Control de soldadura", "Laboratorio", 8, "Resistencia de la soldadura (arrancamiento)", "", "CC", "≥ 18 kN; límite de aviso 20 kN", "Máquina de tracción", "3", "Por lote", "Registro de laboratorio", "< 20 kN: 5 piezas + macro. < 18 kN: bloquear lote y abrir 8D"],
             [40, "Control de soldadura", "Laboratorio", 9, "Penetración en la raíz W2", "", "CC", "≥ 1,0 mm", "Macrografía", "1", "Semanal por programa", "Registro de laboratorio", "Bloquear producción de la semana; análisis"],
             [40, "Control de soldadura", "Puesto de cincel", 10, "Integridad del cordón", "", "CC", "Sin separación", "Ensayo de cincel", "1", "Por turno", "Hoja de control L2", "Bloquear desde último OK"],
             [50, "Cataforesis", "EC-01", 11, "Espesor e-coat", "", "SC", "18–25 µm", "Medidor magnético", "1/20", "Por bastidor", "Registro L3", "Recubrir o segregar"],
             [60, "Inspección final", "VI-01", 12, "Identificación LH/RH", "", "", "Forma según referencia", "Visión artificial", "100 %", "Continuo", "Registro VI-01", "Segregar"],
         ]),
    dict(code="PC-AR1007-01", version="v3", refs=["AR-1007"], title="Plan de control — Brida de anclaje del cinturón",
         header=[("Plan de control", "PC-AR1007-01 · v3 · Serie"), ("Pieza", "AR-1007 (Leitzaran Seating)"), ("AMFE asociado", "AMFE-AR1007-01 v3"),
                 ("Fecha", "30/09/2025"), ("Aprobado", "I. Etxeberria (Calidad)")],
         rows=[
             [10, "Recepción de bobina", "—", 1, "Rm", "", "", "560–700 MPa; AR-1007 ≤ 650 MPa", "Certificado 3.1", "1", "Cada bobina", "Registro de recepción", "Bloquear bobina > 650 MPa"],
             [20, "Punzonado P3", "PR-250 / MT-07", 2, "", "Golpes desde el último afilado", "", "≤ 60.000 (PROC-L1-AFILADO-01 v3)", "Contador de prensa", "—", "Continuo (alarma)", "Registro de afilados MT-07", "Retirar matriz; montar punzón de repuesto"],
             [20, "Punzonado P3", "PR-250 / MT-07", 3, "Ø agujero de anclaje", "", "CC", "Ø11,20 +0,15/0", "Calibre tampón pasa/no pasa", "1", "Primera pieza y cada 2 h", "Hoja de control L1", "Parar; segregar desde último OK; aviso a matricería"],
             [20, "Punzonado P3", "PR-250 / MT-07", 4, "Rebaba en el agujero", "", "CC", "≤ 0,2 mm", "Reloj comparador, 4 posiciones", "1", "Primera pieza y cada 2 h", "Hoja de control L1", "Parar; segregar desde último OK; afilar"],
             [20, "Punzonado P3", "Calibre tampón", 5, "", "Desgaste del calibre", "", "Ø11,200 ± 0,003", "Micrómetro", "—", "Semanal; calibración cada 3 meses", "Registro de metrología", "Retirar calibre; revisar lotes desde la última verificación"],
             [30, "Doblado P5", "PR-250 / MT-07", 6, "Ángulo de alas", "", "", "90 ± 0,5°", "Útil de control", "1", "Por turno", "Hoja de control L1", "Ajustar reestampado"],
             [40, "Inspección final", "VI-01", 7, "Agujero libre de pintura", "", "", "Tampón pasa", "Calibre tampón", "1/50", "Continuo", "Registro VI-01", "Segregar y repasar"],
         ]),
    dict(code="PC-AR1006-01", version="v3", refs=["AR-1006", "AR-1010"], title="Plan de control — Soldadura de tuercas (AR-1006 / AR-1010)",
         header=[("Plan de control", "PC-AR1006-01 · v3 · Serie"), ("Piezas", "AR-1006, AR-1010 (OEM Norte)"), ("AMFE asociado", "AMFE-AR1006-01 v3"),
                 ("Fecha", "22/07/2025"), ("Aprobado", "I. Etxeberria (Calidad)")],
         rows=[
             [10, "Recepción de tuercas", "XRF", 1, "Espesor de recubrimiento Zn-Ni", "", "", "8–11 µm (acuerdo S-BIDA)", "XRF", "5", "Cada lote", "Registro de recepción", "Bloquear lote; aviso a Compras"],
             [10, "Alimentación", "AT-01", 2, "", "Presencia y orientación de tuerca", "", "Tuerca M8 presente", "Sensor en electrodo", "100 %", "Continuo", "Autómata", "Parada automática"],
             [20, "Soldadura por proyección", "SP-01 / estación de tuercas CR-02", 3, "", "Alineación de electrodos", "", "≤ 0,2 mm", "Galga de alineación", "—", "Tras cada intervención", "Registro de mantenimiento", "No liberar máquina"],
             [20, "Soldadura por proyección", "SP-01 / estación de tuercas CR-02", 4, "Par de arrancamiento de la tuerca", "", "SC", "≥ 40 Nm", "Llave dinamométrica", "3", "Cada 2 h y 1ª pieza tras mantenimiento", "Hoja de control L2", "Bloquear desde último OK; 100 % par de prueba 40 Nm"],
             [20, "Soldadura por proyección", "Tampón roscado", 5, "Rosca libre", "", "", "M8 pasa", "Tampón roscado", "1", "Cada 2 h", "Hoja de control L2", "Segregar"],
         ]),
    dict(code="PC-AR1002-01", version="v3", refs=["AR-1002", "AR-1005"], title="Plan de control — Estampación PR-400 (AR-1002 / AR-1005)",
         header=[("Plan de control", "PC-AR1002-01 · v3 · Serie"), ("Piezas", "AR-1002, AR-1005 (Ribera Chassis Systems)"), ("AMFE asociado", "AMFE-AR1002-01 v4"),
                 ("Fecha", "20/06/2025"), ("Aprobado", "I. Etxeberria (Calidad)")],
         rows=[
             [10, "Recepción de bobina", "Laboratorio", 1, "Límite elástico (AR-1005)", "", "CC", "Re ≥ 360 MPa", "Ensayo de tracción", "1", "Cada bobina", "Registro de laboratorio", "Bloquear bobina"],
             [20, "Estampación MT-02", "PR-400", 2, "Planitud cara de apoyo (AR-1002)", "", "SC", "≤ 0,5 mm", "Mármol + galga", "1", "Cada 2 h y tras reparación de matriz", "Hoja de control L1", "Parar; segregar desde último OK"],
             [20, "Estampación MT-05", "PR-400", 3, "Espesor en radio (AR-1005)", "", "CC", "≥ 2,12 mm", "Micrómetro de bola", "1", "Por lote", "Hoja de control L1", "Segregar"],
             [30, "Soldadura MIG CR-02 (AR-1005)", "CR-02", 4, "Penetración", "", "CC", "≥ 1,0 mm", "Macrografía", "1", "Semanal", "Registro de laboratorio", "Bloquear semana"],
             [40, "Ensayo de carga (AR-1005)", "Laboratorio", 5, "Deformación permanente a 12 kN", "", "CC", "≤ 0,5 mm", "Banco de carga", "3", "Por lote", "Registro de laboratorio", "Bloquear lote; 8D"],
             [50, "Embalaje", "Expedición", 6, "Estado de separadores y caja", "", "", "Separador 5 mm íntegro; máx. 4 alturas", "Visual", "100 %", "Cada caja", "IT-L3-VI01-01", "Reembalar"],
         ]),
    dict(code="PC-L3-EC01-01", version="v2", refs=["AR-1009"], title="Plan de control — Cataforesis EC-01",
         header=[("Plan de control", "PC-L3-EC01-01 · v2 · Serie"), ("Alcance", "Todas las referencias; SC de espesor en AR-1009"), ("AMFE asociado", "AMFE-L3-EC01-01 v2"),
                 ("Fecha", "26/08/2024"), ("Aprobado", "I. Etxeberria (Calidad)")],
         rows=[
             [10, "Recepción química S-ARAK", "Laboratorio L3", 1, "", "Fosfatos y conductividad del lote", "", "Fosfatos < 20 ppm tras adición simulada", "Titulación", "1", "Cada lote", "Registro L3", "No añadir; aviso a Compras"],
             [20, "Aclarado DI", "EC-01", 2, "", "Conductividad", "", "< 20 µS/cm", "Conductímetro en línea", "—", "Continuo (alarma)", "Registro continuo", "Parada de línea"],
             [30, "Cataforesis", "EC-01", 3, "Espesor e-coat", "", "SC", "18–25 µm", "Medidor magnético 5 puntos", "1/20", "Por bastidor", "Registro L3", "Segregar"],
             [30, "Cataforesis", "EC-01", 4, "Poros / cráteres", "", "SC", "Sin poros visibles", "Visual 1.000 lux", "5", "Por bastidor", "Registro L3", "Segregar bastidor; revisar baño"],
             [40, "Curado", "HC-01", 5, "", "Temperatura / tiempo", "", "180 °C / 20 min", "Registrador", "—", "Continuo", "Registro HC-01", "Segregar"],
             [50, "Validación", "Laboratorio externo", 6, "Resistencia a niebla salina", "", "SC", "Sin óxido rojo a 480 h", "ISO 9227 NSS", "3", "Mensual", "Informe de laboratorio", "Bloquear producción del mes; 8D"],
         ]),
]


def _pc(spec) -> dict:
    wb = new_workbook()
    ws = wb.active
    ws.title = "Plan de control"
    ws.cell(row=1, column=1, value=spec["title"]).font = Font(bold=True, size=13, color="1F4E79")
    top = xlsx_header_block(ws, spec["header"], top=2)
    write_table(ws, top, PC_HEAD, spec["rows"], PC_W)
    rel = f"calidad/planes-control/{spec['code']}_{spec['version']}.xlsx"
    save_wb(wb, rel)
    return {"doc_id": spec["code"], "version": spec["version"], "title": spec["title"], "path": rel,
            "folder": "calidad/planes-control", "doc_type": "PC", "language": "ES", "part_refs": spec["refs"], "pages": None,
            "sheets": ["Plan de control"]}


def build_all() -> list[dict]:
    out = [_amfe_clean(s) for s in (AMFE_AR1003, AMFE_AR1007, AMFE_AR1002, AMFE_EC01)]
    out.append(_amfe_dirty())
    out += [_pc(s) for s in PCS]
    return out
