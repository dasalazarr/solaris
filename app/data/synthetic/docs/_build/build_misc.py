"""Mantenimiento, evaluaciones de proveedor, costes de no calidad y plantillas 8D. DATOS SINTÉTICOS (ADR-0004)."""
from __future__ import annotations

from datetime import date

from openpyxl.styles import Font

import erp_seed as erp
import maint
from common import Doc, fdate, fnum, new_workbook, save_wb, write_table, xlsx_header_block


# ----------------------------------------------------------------------------- mantenimiento
def plan_afilado() -> dict:
    wb = new_workbook(creator="Mantenimiento Arga")
    ws = wb.active
    ws.title = "Plan"
    ws["A1"] = "Plan de afilado y rectificado de matrices L1"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    top = xlsx_header_block(ws, [("Código", "PROC-L1-AFILADO-01 · versión v3"), ("Fecha", "20/01/2025"),
                                 ("Elaborado", "K. Arbizu (Mantenimiento)"), ("Aprobado", "K. Arbizu (Mantenimiento) · J. Mendia (Dirección)"),
                                 ("Subcontrata", "S-ZIDA Tratamientos Zidacos S.L. (afilado y rectificado)")], top=2)
    rows = []
    for p in erp.load()["parts"]:
        itv = {"MT-07": 60000, "MT-12": 50000, "MT-01": 60000, "MT-08": 60000}.get(p["die"], 80000 if p["press"] == "PR-400" else 60000)
        crit = "Punzón P3 Ø11,20 (CC)" if p["die"] == "MT-07" else ("Punzones de taladro" if p["die"] in ("MT-01", "MT-08") else "Aristas de corte")
        rows.append([p["die"], p["press"], p["ref"], crit, itv, "S-ZIDA", "5 días laborables", "Sí" if p["die"] in ("MT-07", "MT-01") else "No",
                     "Alarma de contador en prensa al 95 %"])
    write_table(ws, top, ["Matriz", "Prensa", "Referencia", "Elemento crítico", "Intervalo de afilado (golpes)", "Proveedor",
                          "Plazo máx.", "Punzón de repuesto", "Control"], rows, [8, 8, 10, 26, 16, 10, 14, 12, 30])
    ws2 = wb.create_sheet("Control de cambios")
    write_table(ws2, 1, ["Versión", "Fecha", "Cambio", "Motivo", "Aprobado"], [
        ["v1", "15/02/2023", "Emisión del plan. MT-07: 40.000 golpes", "Plan inicial (datos del matricero)", "K. Arbizu"],
        ["v2", "20/11/2024", "Plazo máximo S-ZIDA 5 días laborables; punzón de repuesto para MT-07 y MT-01; alarma de contador", "Reclamación de cliente por rebaba en AR-1007", "K. Arbizu"],
        ["v3", "20/01/2025", "MT-07: intervalo de afilado de 40.000 a 60.000 golpes. MT-01 y MT-08: de 50.000 a 60.000 golpes",
         "Propuesta de Compras y S-ZIDA (nuevo proceso de rectificado). Ahorro estimado 7.800 €/año en afilados", "K. Arbizu · J. Mendia"],
    ], [8, 12, 60, 55, 18])
    ws3 = wb.create_sheet("Registro MT-07")
    ev, per_lot = maint.mt07_history()
    last_lot = erp.lots_of("AR-1007")[-1]
    rows3 = [[fdate(e["date"]), e["ot"], "MT-07 / punzón P3", fnum(e["strokes"]), fnum(e["interval"]),
              fnum(e["strokes"] - e["interval"]) if e["strokes"] > e["interval"] else "—", "S-ZIDA", e["reason"]] for e in ev]
    write_table(ws3, 1, ["Fecha", "OT", "Elemento", "Golpes al afilar", "Intervalo vigente", "Exceso", "Proveedor", "Tipo / observaciones"],
                rows3, [12, 14, 18, 16, 16, 10, 10, 50])
    r = len(rows3) + 3
    ws3.cell(row=r, column=1, value=f"Estado a {fdate(last_lot['production_date'])}: {fnum(per_lot[last_lot['lot_code']])} golpes desde el último afilado. "
                                    "Afilado pedido a S-ZIDA (OT-MT-26-006), programado para el 29/09/2026.").font = Font(italic=True)
    rel = "produccion/mantenimiento/PROC-L1-AFILADO-01_v3.xlsx"
    save_wb(wb, rel)
    refs = sorted({p["ref"] for p in erp.load()["parts"]})
    return {"doc_id": "PROC-L1-AFILADO-01", "version": "v3", "title": "Plan de afilado y rectificado de matrices L1 (incluye registro MT-07)",
            "path": rel, "folder": "produccion/mantenimiento", "doc_type": "PROC", "language": "ES", "part_refs": refs, "pages": None,
            "sheets": ["Plan", "Control de cambios", "Registro MT-07"]}


def registro_boquillas() -> dict:
    wb = new_workbook(creator="Producción L2")
    ws = wb.active
    ws.title = "Registro"
    ws["A1"] = "Registro de cambio de boquilla, tubo de contacto y bobina — célula CR-01"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E79")
    top = xlsx_header_block(ws, [("Código", "REG-L2-CR01-01 · v1"), ("Referencia", "IT-L2-CR01-03 §5.2 y §5.3"),
                                 ("Inicio del registro", fdate(maint.LOG_START)), ("Responsable", "Jefe de turno (firma)")], top=2)
    rows = [[fdate(r[0]), *r[1:]] for r in maint.cr01_log()]
    write_table(ws, top, ["Fecha", "Turno", "Hora", "Célula", "Elemento", "Motivo", "Operario", "Firma jefe de turno", "Observaciones"],
                rows, [11, 9, 7, 7, 24, 15, 13, 17, 48])
    ws.freeze_panes = ws.cell(row=top + 1, column=1)
    rel = "produccion/mantenimiento/REG-L2-CR01-01_v1.xlsx"
    save_wb(wb, rel)
    return {"doc_id": "REG-L2-CR01-01", "version": "v1", "title": "Registro de cambio de boquillas, tubo de contacto y bobina en CR-01",
            "path": rel, "folder": "produccion/mantenimiento", "doc_type": "REG", "language": "ES", "part_refs": ["AR-1003", "AR-1004"],
            "pages": None, "sheets": ["Registro"], "rows": len(rows)}


# ----------------------------------------------------------------------------- proveedores
def _eval(code, sup, version, when, period, scores, notes, actions, rating) -> dict:
    s = erp.by_key("suppliers", "code")[sup]
    d = Doc(code, version, f"Evaluación de proveedor {sup}", header_right="Compras — CONFIDENCIAL")
    d.title(f"EVALUACIÓN DE PROVEEDOR — {s['name']} ({sup})", f"Periodo evaluado: {period}")
    d.form([("Código", code), ("Versión", version), ("Proveedor", f"{s['name']} ({sup})"), ("Suministro", s["supplies"]),
            ("Fecha", when), ("Evaluador", "L. Sarasola (Compras) · M. Irigoyen (Calidad)")])
    d.h("1. Puntuación")
    d.table(["Criterio", "Peso", "Resultado", "Puntos"], scores, font=8)
    d.p(f"Clasificación: {rating}", bold=True)
    d.h("2. Observaciones")
    d.bullets(notes)
    d.h("3. Acciones acordadas")
    d.table(["Acción", "Responsable", "Fecha", "Estado"], actions, font=8)
    d.signatures([("Compras", "Leire Sarasola", when, "firmado"), ("Calidad", "Iñaki Etxeberria", when, "firmado")])
    rel = f"compras/proveedores/{code}_{version}.docx"
    d.save(rel)
    return {"doc_id": code, "version": version, "title": f"Evaluación de proveedor {s['name']} ({sup}) — {period}", "path": rel,
            "folder": "compras/proveedores", "doc_type": "EVAL", "language": "ES", "part_refs": [], "pages": None, "supplier": sup}


def evaluaciones() -> list[dict]:
    goie_lot = erp.by_key("material_lots", "lot_code")["S-GOIE-260117"]
    return [
        _eval("EVAL-SGOIE-01", "S-GOIE", "v3", "10/09/2026", "01/2026 – 08/2026",
              [["Calidad (ppm, incidencias)", "40 %", "0 ppm; 0 reclamaciones imputadas", "38"],
               ["Entregas (OTD)", "30 %", "97,2 %", "27"], ["Gestión de cambios", "15 %", "1 PCN notificada con 2 semanas de antelación", "12"],
               ["Servicio / respuesta", "15 %", "Correcto", "13"]],
              ["Hilo MIG G3Si1 Ø1,0 mm y gas M21. Certificados 3.1 conformes en todos los lotes del periodo.",
               f"PCN-GOIE-2026-04 recibida el 05/08/2026: cambio de fabricante de origen del alambrón para el hilo G3Si1 (mismo grado y norma). "
               f"Aceptada por Compras con el certificado del nuevo origen, sin ensayos de soldabilidad en Arga. Primer lote con el nuevo origen: "
               f"{goie_lot['lot_code']} (recepción {fdate(goie_lot['received_date'])}, {fnum(int(goie_lot['qty']))} kg).",
               "En las reclamaciones de soldadura de años anteriores el hilo se revisó y se descartó como causa (certificado conforme).",
               "Proveedor único de hilo MIG: riesgo de suministro medio."],
              [["Solicitar al proveedor datos comparativos de soldabilidad del nuevo origen", "L. Sarasola", "30/09/2026", "Abierta"],
               ["Evaluar segundo proveedor de hilo", "L. Sarasola", "31/12/2026", "Abierta"]], "A (90/100)"),
        _eval("EVAL-SZIDA-01", "S-ZIDA", "v2", "28/01/2026", "2025",
              [["Calidad del afilado", "40 %", "1 rechazo por holgura (03/2025)", "34"], ["Plazo (≤ 5 días laborables)", "35 %", "92 % en plazo", "30"],
               ["Coste", "15 %", "Reducción por ampliación de intervalos", "15"], ["Servicio", "10 %", "Correcto", "8"]],
              ["Afilado y rectificado de matrices de PR-250 y PR-400.",
               "En 01/2025 se aceptó la propuesta de S-ZIDA de ampliar el intervalo de afilado de MT-07 de 40.000 a 60.000 golpes con su nuevo proceso de rectificado (PROC-L1-AFILADO-01 v3). Ahorro 2025: 7.800 €.",
               "Plazo mejorado tras el acuerdo de noviembre de 2024 (máx. 5 días laborables)."],
              [["Revisión anual de intervalos de afilado con datos de desgaste", "K. Arbizu", "30/06/2026", "Pendiente"]], "B (87/100)"),
        _eval("EVAL-SBIDA-01", "S-BIDA", "v2", "30/01/2026", "2025",
              [["Calidad", "40 %", "1 incidencia (recubrimiento Zn-Ni en el límite superior, 05/2025)", "28"], ["Entregas", "30 %", "98,5 %", "29"],
               ["Gestión de cambios", "15 %", "PCN de ventana de recubrimiento 8–11 µm emitida", "13"], ["Servicio", "15 %", "Correcto", "13"]],
              ["Tuercas soldables M8 Zn-Ni para AR-1006 y AR-1010.", "Incidencia de 2025: lote S-BIDA-250129 con espesor de recubrimiento en el límite superior; tuercas que giran en cliente.",
               "Ventana restringida 8–11 µm acordada desde 07/2025; XRF en recepción en Arga."],
              [["Auditoría de proceso en S-BIDA (baño Zn-Ni)", "L. Sarasola", "31/05/2026", "Planificada"]], "B (83/100)"),
    ]


# ----------------------------------------------------------------------------- dirección (confidencial)
COST_2025 = [  # 8D, clasificación de costes (€): selección, chatarra, transporte urgente, cargo del cliente, horas internas
    ("8D-ARGA-2025-002", 6800, 1900, 1200, 12500, 3400), ("8D-ARGA-2025-005", 2100, 1400, 600, 0, 2600),
    ("8D-ARGA-2025-006", 3900, 1100, 400, 2500, 1800), ("8D-ARGA-2025-008", 4200, 900, 800, 8000, 2900),
    ("8D-ARGA-2025-010", 5600, 2300, 700, 6000, 2400), ("8D-ARGA-2025-011", 3300, 700, 900, 5000, 1700),
    ("8D-ARGA-2025-014", 7400, 800, 1500, 15000, 3800), ("8D-ARGA-2025-015", 4900, 450, 300, 3500, 1500),
]


def costes() -> dict:
    wb = new_workbook(creator="Dirección")
    ws = wb.active
    ws.title = "CNC 2025"
    ws["A1"] = "Costes de no calidad externos 2025 — CONFIDENCIAL (Dirección)"
    ws["A1"].font = Font(bold=True, size=13, color="C00000")
    top = xlsx_header_block(ws, [("Código", "COST-DIR-01 · v1"), ("Fecha", "31/01/2026"), ("Elaborado", "Control de gestión"),
                                 ("Distribución", "Dirección general (J. Mendia). No difundir.")], top=2)
    comp = {c["report_8d_id"]: c for c in erp.load()["complaints"] if c["report_8d_id"]}
    rows = []
    for rid, *vals in COST_2025:
        c = comp[rid]
        rows.append([rid, c["customer_code"], c["part_ref"], *vals, sum(vals)])
    total = [sum(r[i] for r in rows) for i in range(3, 9)]
    rows.append(["TOTAL", "", "", *total])
    write_table(ws, top, ["8D", "Cliente", "Ref.", "Selección (€)", "Chatarra (€)", "Transporte urgente (€)", "Cargo del cliente (€)",
                          "Horas internas (€)", "Total (€)"], rows, [18, 9, 9, 13, 12, 18, 18, 16, 12])
    ws2 = wb.create_sheet("Margen por cliente")
    write_table(ws2, 1, ["Cliente", "Ventas 2025 (k€)", "Margen bruto (%)", "CNC externo (k€)", "CNC / ventas (%)"], [
        ["C-OEMN", 14820, 17.8, round(sum(r[8] for r in rows[:-1] if r[1] == "C-OEMN") / 1000, 1), None],
        ["C-RIBE", 6240, 14.1, round(sum(r[8] for r in rows[:-1] if r[1] == "C-RIBE") / 1000, 1), None],
        ["C-LEIZ", 4710, 15.6, round(sum(r[8] for r in rows[:-1] if r[1] == "C-LEIZ") / 1000, 1), None],
    ], [10, 16, 16, 16, 16])
    for r in range(2, 5):
        ws2.cell(row=r, column=5, value=f"=ROUND(D{r}/B{r}*100,2)")
    ws2["A7"] = "Nota: los cargos de OEM Norte incluyen gastos de gestión de reclamaciones en características CC."
    rel = "direccion/confidencial/COST-DIR-01_v1.xlsx"
    save_wb(wb, rel)
    return {"doc_id": "COST-DIR-01", "version": "v1", "title": "Costes de no calidad externos 2025 y margen por cliente", "path": rel,
            "folder": "direccion/confidencial", "doc_type": "COST", "language": "ES", "part_refs": [], "pages": None,
            "sheets": ["CNC 2025", "Margen por cliente"]}


# ----------------------------------------------------------------------------- plantillas
def tpl_arga() -> dict:
    d = Doc("TPL-8D-ARGA", "v2", "Plantilla de informe 8D", header_right="Plantilla interna")
    d.title("INFORME 8D / 8D REPORT — TPL-8D-ARGA v2", "Plantilla interna de Componentes Arga S.L. Usar en ES o EN según el idioma del cliente.")
    d.form([("Cliente / Customer", ""), ("Nº reclamación / Complaint no.", ""), ("Referencia / Part", ""), ("Lote / Lot", ""),
            ("Fecha de recepción / Date received", ""), ("Cantidad / Qty", ""), ("Línea / Line", ""), ("Nº 8D", "8D-ARGA-AAAA-nnn")])
    guide = [
        ("D1 · Equipo / Team", "Champion (Responsable de Calidad), líder del 8D y personas de Producción, Ingeniería, Mantenimiento y Compras según el caso. Incluir el turno afectado."),
        ("D2 · Descripción del problema (5W2H) / Problem description", "Qué, dónde, cuándo, quién, cuál, cómo y cuántas. Lote, fecha de fabricación, turno, equipo y lotes de material (datos del ERP)."),
        ("D3 · Contención / Containment", "Plazo: 48 h (Tier 1) o el que exija el cliente. Tabla de lotes sospechosos: stock Arga, tránsito y cliente; resultados de la selección; punto limpio."),
        ("D4 · Causa raíz / Root cause", "Ishikawa 6M + 5 porqués de ocurrencia y de no detección. Verificar la causa reproduciendo el defecto. Revisar 8D anteriores de la misma pieza o defecto."),
        ("D5 · Acciones correctivas / Corrective actions", "Una acción por causa confirmada, con responsable y fecha."),
        ("D6 · Implantación y validación / Implementation", "Evidencias de implantación y resultados (datos, no opiniones)."),
        ("D7 · Prevenir la repetición / Prevent recurrence", "Actualizar AMFE, plan de control e IT. Extender a piezas y equipos similares. Verificar en TODOS los turnos."),
        ("D8 · Cierre / Closure", "Cierre tras comprobar la eficacia (mín. 60 días) y aceptación del cliente."),
    ]
    for h, g in guide:
        d.h(h, 1)
        d.p(g, italic=True, size=9)
        d.table(None, [[""], [""]], font=9)
    rel = "calidad/plantillas/TPL-8D-ARGA_v2.docx"
    d.save(rel)
    return {"doc_id": "TPL-8D-ARGA", "version": "v2", "title": "Plantilla interna de informe 8D (ES/EN)", "path": rel,
            "folder": "calidad/plantillas", "doc_type": "TPL", "language": "ES", "part_refs": [], "pages": None}


def tpl_oemn() -> dict:
    d = Doc("TPL-8D-OEMN", "v3", "8D Report", lang="EN", header_left="OEM Norte 8D Report v3", header_right="Supplier Quality form OEMN-SQ-8D",
            accent="3A3A3A")
    d.title("OEM Norte 8D Report v3", "Supplier Quality form OEMN-SQ-8D rev. 3 — fictitious OEM, synthetic demo data")
    d.p("Instructions to suppliers", bold=True)
    d.bullets(["Containment (D3) must be reported within 24 hours of notification.",
               "The complete report (D1–D8) must be returned within 10 working days. D5–D8 may be updated later with evidence.",
               "Use one report per complaint number. Reference previous reports on the same part number where relevant.",
               "Root cause must address both occurrence and non-detection, supported by a verification test.",
               "Attach evidence: containment table by lot, 5 Whys, Ishikawa, updated PFMEA and control plan pages."])
    d.form([("Supplier", ""), ("Supplier code", ""), ("OEM Norte complaint no.", ""), ("Supplier 8D no.", ""), ("Part number", ""),
            ("Part description", ""), ("Special characteristic", ""), ("Supplier lot", ""), ("Complaint date", ""), ("Qty complained", ""),
            ("Containment due (24 h)", ""), ("8D due (10 WD)", ""), ("Report status", "OPEN / CLOSED"), ("Closure date", "")])
    sections = [
        ("D1 · Team", ["Name", "Role", "Department"]),
        ("D2 · Problem description (5W2H)", ["5W2H", "Description"]),
        ("D3 · Interim containment actions", ["Lot", "Part", "Prod. date", "Shift", "Qty OK produced", "Shipped", "Supplier stock", "Action", "Result"]),
        ("D4 · Root cause analysis", ["Category (6M)", "Potential cause", "Verification / status"]),
        ("D5 · Permanent corrective actions (chosen)", ["Action", "Type", "Owner", "Due date"]),
        ("D6 · Implementation and validation", ["Evidence", "Result"]),
        ("D7 · Actions to prevent recurrence", ["Action", "Document / system", "Owner", "Due date", "Status"]),
        ("D8 · Closure and team recognition", ["Role", "Name", "Date", "Signature"]),
    ]
    for h, cols in sections:
        d.h(h, 1)
        d.table(cols, [[""] * len(cols), [""] * len(cols)], font=8)
    rel = "calidad/plantillas/TPL-8D-OEMN_v3.docx"
    d.save(rel)
    return {"doc_id": "TPL-8D-OEMN", "version": "v3", "title": "OEM Norte 8D Report v3 (plantilla del cliente, ficticia)", "path": rel,
            "folder": "calidad/plantillas", "doc_type": "TPL", "language": "EN", "part_refs": [], "pages": None}


def build_all() -> list[dict]:
    return [plan_afilado(), registro_boquillas(), *evaluaciones(), costes(), tpl_arga(), tpl_oemn()]
